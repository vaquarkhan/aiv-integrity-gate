# Copyright 2026 Viquar Khan
# SPDX-License-Identifier: Apache-2.0
"""Validate LinkedIn carousel PPTX/PDF/renders and create a contact sheet."""
from pathlib import Path
from pptx import Presentation
from PIL import Image, ImageDraw
import fitz
import re
import sys
import zipfile

HERE=Path(__file__).resolve().parent
PPT=HERE/'AIV-LinkedIn-Carousel.pptx'
PDF=HERE/'AIV-LinkedIn-Carousel.pdf'
RENDER=HERE/'rendered-linkedin-carousel'
CONTACT=HERE/'AIV-LinkedIn-Carousel-contact-sheet.png'
errors=[]
prs=Presentation(PPT)
if len(prs.slides)!=16: errors.append(f'PPT pages {len(prs.slides)} != 16')

expected=[
 'AI agents create change faster',
 'Why existing solutions do not close the gap',
 'AIV turns change into verifiable evidence',
 'How AIV works end to end',
 '16 gates. Four evidence jobs.',
 'Stop guaranteed breakage first',
]
page_texts=[]
all_text=[]
for n,slide in enumerate(prs.slides,1):
    page=[]
    for sh in slide.shapes:
        if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width or sh.top+sh.height>prs.slide_height:
            errors.append(f'page {n}: shape outside bounds')
        if hasattr(sh,'text') and sh.text.strip():
            page.append(sh.text.strip()); all_text.append(sh.text)
    page_texts.append(page)
for i,title in enumerate(expected):
    if title not in page_texts[i]: errors.append(f'page {i+1}: missing opening title {title!r}')
text='\n'.join(all_text)
for token in ['syntax','invariant','dependency','package-existence','test-adequacy','security','design','rules','density','cohesion','duplication','external','doc-integrity','pr-metadata','license-header','jev','AI-DLC','844','Viquar Khan']:
    if token.lower() not in text.lower(): errors.append(f'missing topic {token}')
if '\u2014' in text or '\u2013' in text: errors.append('contains em/en dash')

with zipfile.ZipFile(PPT) as z:
    core=z.read('docProps/core.xml').decode('utf-8')
    if 'Viquar Khan' not in core: errors.append('PPT author metadata missing')

pdf=fitz.open(PDF)
if pdf.page_count!=16: errors.append(f'PDF pages {pdf.page_count} != 16')
for i,page in enumerate(pdf):
    ratio=page.rect.width/page.rect.height
    if abs(ratio-0.8)>0.01: errors.append(f'PDF page {i+1}: ratio {ratio:.3f} != 0.8')
pdf.close()

pages=sorted(RENDER.glob('Slide*.PNG'),key=lambda p:int(re.search(r'\d+',p.stem).group()))
if len(pages)!=16: errors.append(f'rendered pages {len(pages)} != 16')
sheet=Image.new('RGB',(1080,4*365),'#E2E8F0'); draw=ImageDraw.Draw(sheet)
for i,path in enumerate(pages):
    with Image.open(path) as im:
        if im.size!=(1080,1350): errors.append(f'{path.name}: {im.size} != 1080x1350')
        thumb=im.convert('RGB'); thumb.thumbnail((260,325))
        x=(i%4)*270+5; y=(i//4)*365+5
        sheet.paste(thumb,(x,y)); draw.text((x+4,y+330),f'{i+1:02d}',fill='#0B132B')
sheet.save(CONTACT)

if errors:
    print('LinkedIn carousel validation failed:')
    for e in errors: print(' -',e)
    sys.exit(1)
print(f'LinkedIn carousel validation OK: 16 PPT/PDF/rendered pages; {CONTACT.name}')
