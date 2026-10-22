# Copyright 2026 Viquar Khan
# SPDX-License-Identifier: Apache-2.0
"""Validate the generated sales deck and build a rendered contact sheet."""
from pathlib import Path
from pptx import Presentation
from PIL import Image, ImageDraw
import zipfile
import re
import sys

HERE = Path(__file__).resolve().parent
DECK = HERE / "AIV-Integrity-Gate-Sales-Deck.pptx"
RENDERED = HERE / "rendered-sales-deck-v2"
CONTACT = HERE / "AIV-Integrity-Gate-Sales-Deck-contact-sheet.png"
errors = []

prs = Presentation(DECK)
if len(prs.slides) != 21:
    errors.append(f"expected 21 slides, found {len(prs.slides)}")
expected_opening = [
    "AI agents create change faster",
    "Why existing solutions do not close the gap",
    "AIV turns untrusted change into verifiable evidence",
    "What is AIV? Complete product architecture",
    "What gates exist, and where should each be used?",
    "Where teams use AIV",
]
actual_opening = []
for slide in list(prs.slides)[:6]:
    actual_opening.append(next((shape.text.strip() for shape in slide.shapes
                                if hasattr(shape, "text") and shape.text.strip()), ""))
for index, expected in enumerate(expected_opening):
    if actual_opening[index] != expected:
        errors.append(f"slide {index+1}: expected opening title '{expected}', found '{actual_opening[index]}'")

required = ["syntax", "invariant", "dependency", "package-existence", "test-adequacy", "security",
            "design", "rules", "density", "cohesion", "duplication", "external", "doc-integrity",
            "pr-metadata", "license-header", "jev", "AI-DLC", "aiv verify", "Audit", "Pilot"]
all_text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text_frame"))
for token in required:
    if token.lower() not in all_text.lower():
        errors.append(f"missing required topic: {token}")
if "\u2014" in all_text or "\u2013" in all_text:
    errors.append("deck contains em/en dash")
for number, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
            errors.append(f"slide {number}: shape outside slide bounds")

with zipfile.ZipFile(DECK) as archive:
    slide_xml = [n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
    if len(slide_xml) != 21:
        errors.append("pptx archive slide count mismatch")
    core = archive.read("docProps/core.xml").decode("utf-8")
    if "Viquar Khan" not in core:
        errors.append("author metadata missing")

slides = sorted(RENDERED.glob("Slide*.PNG"), key=lambda p: int(re.search(r"\d+", p.stem).group()))
if len(slides) != 21:
    errors.append(f"expected 21 rendered slides, found {len(slides)}")
else:
    thumb_w, thumb_h = 400, 225
    cols, rows = 4, 6
    sheet = Image.new("RGB", (cols * thumb_w, rows * (thumb_h + 28)), "#E2E8F0")
    draw = ImageDraw.Draw(sheet)
    for i, path in enumerate(slides):
        with Image.open(path) as image:
            if image.size != (1600, 900):
                errors.append(f"{path.name}: unexpected size {image.size}")
            image = image.convert("RGB")
            image.thumbnail((thumb_w, thumb_h))
            x = (i % cols) * thumb_w
            y = (i // cols) * (thumb_h + 28)
            sheet.paste(image, (x, y))
            draw.text((x + 8, y + thumb_h + 5), f"{i+1:02d}", fill="#0B132B")
    sheet.save(CONTACT)

if errors:
    print("Deck validation failed:")
    for error in errors:
        print(" -", error)
    sys.exit(1)
print(f"Deck validation OK: {len(prs.slides)} slides; contact sheet {CONTACT.name}")
