# Copyright 2026 Viquar Khan
# SPDX-License-Identifier: Apache-2.0
"""Build a self-contained 4:5 LinkedIn PDF carousel as editable PowerPoint."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

HERE = Path(__file__).resolve().parent
OUT = HERE / "AIV-LinkedIn-Carousel.pptx"
W, H = 8.0, 10.0
NAVY = RGBColor(11, 19, 43)
INK = RGBColor(23, 37, 61)
SLATE = RGBColor(71, 85, 105)
MUTED = RGBColor(100, 116, 139)
TEAL = RGBColor(15, 118, 110)
MINT = RGBColor(223, 247, 242)
BLUE = RGBColor(232, 241, 255)
AMBER = RGBColor(255, 244, 214)
ROSE = RGBColor(254, 236, 236)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(203, 213, 225)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


def shape(slide, x, y, w, h, fill=WHITE, line=LINE, rounded=True):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line
    return s


def txt(slide, value, x, y, w, h, size=20, color=INK, bold=False,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font="Aptos"):
    b = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = value; r.font.name = font; r.font.size = Pt(size)
    r.font.bold = bold; r.font.color.rgb = color
    return b


def base(page, dark=False):
    s = prs.slides.add_slide(blank)
    bg = NAVY if dark else WHITE
    s.background.fill.solid(); s.background.fill.fore_color.rgb = bg
    if not dark:
        bar = shape(s, .45, .42, .65, .06, TEAL, TEAL, False)
    txt(s, "AIV INTEGRITY GATE", .5, 9.55, 3.7, .18, 8.5, RGBColor(148,163,184) if dark else MUTED, True)
    txt(s, f"{page:02d} / 16", 6.7, 9.55, .8, .18, 8.5, RGBColor(148,163,184) if dark else MUTED, True, PP_ALIGN.RIGHT)
    return s


def heading(slide, title, kicker=None):
    if kicker: txt(slide, kicker.upper(), .5, .65, 7.0, .25, 10, TEAL, True)
    txt(slide, title, .5, 1.02, 7.0, 1.45, 29, NAVY, True)


def card(slide, y, title, body, fill=BLUE, accent=TEAL, h=1.32):
    shape(slide, .55, y, 6.9, h, fill, LINE)
    shape(slide, .55, y, .08, h, accent, accent, False)
    txt(slide, title, .82, y+.18, 6.2, .32, 16, NAVY, True)
    txt(slide, body, .82, y+.58, 6.15, h-.68, 14, SLATE)


def mini(slide, x, y, w, h, title, body, fill=BLUE, accent=TEAL):
    shape(slide,x,y,w,h,fill,LINE)
    shape(slide,x,y,.07,h,accent,accent,False)
    txt(slide,title,x+.2,y+.15,w-.35,.34,14,NAVY,True)
    txt(slide,body,x+.2,y+.55,w-.35,h-.68,13,SLATE)


def bullets(slide, lines, x, y, w, h, size=15, color=INK):
    b = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=b.text_frame; tf.clear(); tf.word_wrap=True
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text="•  "+line; p.font.name="Aptos"; p.font.size=Pt(size); p.font.color.rgb=color; p.space_after=Pt(8)
    return b


def arrow(slide,x1,y1,x2,y2,color=TEAL):
    c=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(2.2); c.line.end_arrowhead=True
    return c


# 1 Hook
s=base(1,True)
txt(s,"AI agents create change faster",.55,.85,6.9,1.0,34,WHITE,True)
txt(s,"than teams can verify it.",.55,1.92,6.9,1.0,34,RGBColor(45,212,191),True)
txt(s,"The bottleneck moved from typing code to proving that a change is safe, complete, and worth human attention.",.58,3.20,6.7,1.2,18,RGBColor(203,213,225))
for i,(n,l) in enumerate([("SECONDS","to generate a large patch"),("HOURS","to understand and verify it"),("ONE GAP","objective evidence before review")]):
    y=5.0+i*1.05; shape(s,.65,y,6.6,.82,RGBColor(20,31,61),RGBColor(51,65,85))
    txt(s,n,.9,y+.20,1.35,.3,14,RGBColor(45,212,191),True)
    txt(s,l,2.35,y+.20,4.55,.35,13,WHITE,True)
txt(s,"Swipe to see how AIV turns an untrusted change into a verifiable decision.",.7,8.55,6.6,.55,14,WHITE,True,PP_ALIGN.CENTER)

# 2 Existing solutions
s=base(2); heading(s,"Why existing solutions do not close the gap","The missing layer")
card(s,2.45,"Compiler / SAST / secret scanner","Excellent specialists, but they run as separate jobs and reports. They do not create one protected evidence receipt for the change.",BLUE,RGBColor(59,130,246),1.35)
card(s,4.02,"LLM reviewer","Useful for contextual judgement, but its answer can vary with prompt, model, context, and time. That is not enough for the only hard gate.",ROSE,RGBColor(220,38,38),1.35)
card(s,5.59,"SDLC orchestrator / policy engine","They sequence work or decide allow/deny. They still need trustworthy code facts as input.",MINT,TEAL,1.35)
shape(s,.75,7.50,6.5,1.05,AMBER,RGBColor(245,158,11))
txt(s,"The gap",1.0,7.70,1.2,.3,16,RGBColor(180,83,9),True)
txt(s,"One diff-scoped, policy-bound evidence receipt that composes these tools.",2.05,7.67,4.85,.55,15,NAVY,True)

# 3 Solution
s=base(3); heading(s,"AIV turns change into verifiable evidence","The solution")
steps=[("1","Capture the proposed diff"),("2","Pin one trusted policy commit"),("3","Freeze every .aiv policy file"),("4","Run objective and advisory gates"),("5","Publish one structured verdict")]
for i,(n,l) in enumerate(steps):
    y=2.45+i*1.12; shape(s,.65,y,6.7,.88,MINT if i in (1,2) else BLUE,LINE)
    shape(s,.82,y+.16,.54,.54,TEAL,TEAL); txt(s,n,.82,y+.26,.54,.2,12,WHITE,True,PP_ALIGN.CENTER)
    txt(s,l,1.62,y+.22,5.35,.35,15,NAVY,True)
txt(s,"AIV does not identify who wrote the code. It asks whether the change is backed by evidence.",.75,8.35,6.5,.62,15,TEAL,True,PP_ALIGN.CENTER)

# 4 how it works
s=base(4); heading(s,"How AIV works end to end","Architecture")
labels=[("INPUT","Git diff or strict JSON",BLUE),("POLICY","Pinned commit + frozen .aiv",MINT),("ENGINE","16 evidence gates",AMBER),("VERDICT","Pass/block + findings",ROSE),("EVIDENCE","JSON/SARIF/GitHub/audit",BLUE)]
for i,(k,l,f) in enumerate(labels):
    y=2.35+i*1.22; shape(s,.75,y,6.5,.9,f,LINE)
    txt(s,k,.98,y+.20,1.1,.25,10,TEAL,True); txt(s,l,2.0,y+.16,4.85,.4,16,NAVY,True)
    if i<4: arrow(s,4.0,y+.92,4.0,y+1.18)
txt(s,"The source workspace and policy workspace stay separate throughout evaluation.",.7,8.65,6.6,.4,13,SLATE,True,PP_ALIGN.CENTER)

# 5 gate map
s=base(5); heading(s,"16 gates. Four evidence jobs.","Gate model")
groups=[("CHANGE INTEGRITY","syntax • invariant\ndependency • package existence",MINT,TEAL),
        ("BEHAVIORAL EVIDENCE","test adequacy • security\ndesign • generic rules",BLUE,RGBColor(59,130,246)),
        ("REVIEW FOCUS","density • cohesion\nduplication • external SARIF",AMBER,RGBColor(245,158,11)),
        ("GOVERNANCE","documentation • PR metadata\nlicense header • advisory Jev",ROSE,RGBColor(220,38,38))]
for i,(h,b,f,a) in enumerate(groups):
    x=.55+(i%2)*3.55; y=2.55+(i//2)*2.45; mini(s,x,y,3.3,1.95,h,b,f,a)
shape(s,.65,7.72,6.7,.92,WHITE,TEAL)
txt(s,"Hard facts can block. Structural or judgement-heavy signals remain advisory.",.9,7.94,6.2,.42,14,TEAL,True,PP_ALIGN.CENTER)

# 6 change integrity
s=base(6); heading(s,"Stop guaranteed breakage first","Gate family 1")
card(s,2.35,"syntax","Strict parse result or visible unverified parser/toolchain state.",MINT,TEAL,1.18)
card(s,3.72,"invariant","Conflict markers, placeholders, agent edit artifacts, assistant chatter, and tautology tests.",MINT,TEAL,1.35)
card(s,5.27,"dependency","Java and Python imports must be backed by manifest, standard-library, or first-party evidence.",MINT,TEAL,1.35)
card(s,6.82,"package-existence","New package declarations are classified as found, absent, or unverified against a registry/mirror.",MINT,TEAL,1.35)
txt(s,"Use: local + CI. Registry access is an explicit external input.",.75,8.55,6.5,.4,13,SLATE,True,PP_ALIGN.CENTER)

# 7 behavioral
s=base(7); heading(s,"Demand evidence around behavior","Gate family 2")
card(s,2.35,"test-adequacy","Added test functions must contain assertion, expectation, or verification evidence.",BLUE,RGBColor(59,130,246),1.25)
card(s,3.80,"security","High-confidence credential formats and keyword-bound secret candidates on added lines.",BLUE,RGBColor(59,130,246),1.25)
card(s,5.25,"design","Repository-owned forbidden/required architecture patterns from frozen YAML policy.",BLUE,RGBColor(59,130,246),1.25)
card(s,6.70,"rules","Project or framework pattern, call-argument, and hollow-function checks in generic YAML.",BLUE,RGBColor(59,130,246),1.25)
txt(s,"Use: CI hard gates after repository calibration. Missing-test evidence remains advisory.",.75,8.48,6.5,.55,13,SLATE,True,PP_ALIGN.CENTER)

# 8 review focus
s=base(8); heading(s,"Focus scarce reviewer attention","Gate family 3")
card(s,2.35,"density","Highlights unusually thin Java logic or low-entropy structure.",AMBER,RGBColor(245,158,11),1.18)
card(s,3.72,"cohesion","Highlights changes spread across many repository areas or files.",AMBER,RGBColor(245,158,11),1.18)
card(s,5.09,"duplication","Finds exact normalized added blocks already present in existing source.",AMBER,RGBColor(245,158,11),1.18)
card(s,6.46,"external","Runs trusted tools and folds their SARIF findings into the same severity-aware receipt.",AMBER,RGBColor(245,158,11),1.35)
txt(s,"Use: advisory review routing. External tools may be hard only when the project already trusts them.",.75,8.42,6.5,.62,13,SLATE,True,PP_ALIGN.CENTER)

# 9 governance
s=base(9); heading(s,"Make delivery policy visible","Gate family 4")
card(s,2.35,"doc-integrity","Checks local links, paths, cross-references, required mentions, and command completeness.",ROSE,RGBColor(220,38,38),1.25)
card(s,3.80,"pr-metadata","Checks title, body, untouched template text, and optional disclosure structure.",ROSE,RGBColor(220,38,38),1.25)
card(s,5.25,"license-header","Checks approved project markers on newly added source files.",ROSE,RGBColor(220,38,38),1.25)
card(s,6.70,"jev","Optional model-backed yes/no questions. Always advisory and declared as an external input.",ROSE,RGBColor(220,38,38),1.25)
txt(s,"Use: enable only the governance surfaces your repository owns and can maintain.",.75,8.48,6.5,.4,13,SLATE,True,PP_ALIGN.CENTER)

# 10 trust
s=base(10); heading(s,"The change cannot edit its own yardstick","Trust boundary")
mini(s,.65,2.35,3.1,1.55,"UNTRUSTED","Proposed source, manifests, tests, docs, and metadata",ROSE,RGBColor(220,38,38))
mini(s,4.25,2.35,3.1,1.55,"TRUSTED","One base ref resolved once to an immutable commit OID",MINT,TEAL)
arrow(s,3.75,3.12,4.20,3.12)
card(s,4.35,"Freeze the complete policy tree","AIV materializes regular .aiv Git blobs into a private snapshot. It rejects drift, ignored extras, symlinks, type changes, and path escapes.",BLUE,RGBColor(59,130,246),1.55)
card(s,6.18,"Bind the decision","Request bytes + canonical config + exact frozen sidecars + policy commit produce the evidence digest.",AMBER,RGBColor(245,158,11),1.35)
txt(s,"The bytes consumed by the gates are the bytes named by the evidence.",.75,8.38,6.5,.55,15,TEAL,True,PP_ALIGN.CENTER)

# 11 outputs
s=base(11); heading(s,"One verdict. Multiple delivery surfaces.","Integration")
items=[("Developer","terminal + exit code"),("CI","JSON + SARIF"),("GitHub","Checks + comments + labels"),("Agent / MCP host","aiv verify schema v2"),("AI-DLC","advisory or fail-closed sensor"),("Governance","audit chain + HTML/JSONL report")]
for i,(h,b) in enumerate(items):
    x=.55+(i%2)*3.55; y=2.35+(i//2)*1.72; mini(s,x,y,3.3,1.35,h,b,BLUE if i%2==0 else MINT,TEAL)
txt(s,"No publisher reruns a gate. Every surface consumes the same AIVResult.",.75,8.15,6.5,.55,14,TEAL,True,PP_ALIGN.CENTER)

# 12 AI-DLC
s=base(12); heading(s,"AIV fits inside an AI-driven lifecycle","AI-DLC")
steps=[("1","Agent produces code or a full diff request"),("2","AI-DLC invokes the AIV sensor"),("3","AIV pins policy and computes evidence"),("4","AI-DLC records, advances, or routes intervention")]
for i,(n,l) in enumerate(steps):
    y=2.45+i*1.25; shape(s,.7,y,6.6,.92,BLUE if i in (0,3) else MINT,LINE)
    txt(s,n,.9,y+.23,.45,.25,13,TEAL,True,PP_ALIGN.CENTER); txt(s,l,1.55,y+.19,5.35,.4,14,NAVY,True)
card(s,7.65,"Two contracts","Advisory mode reports without blocking. Blocking mode requires trusted policy, fails closed, and rejects undeclared external inputs.",AMBER,RGBColor(245,158,11),1.25)

# 13 use cases
s=base(13); heading(s,"Where AIV creates value","Use cases")
card(s,2.35,"Before commit","Catch objective edit/parse/secret/test failures before a pull request exists.",BLUE,RGBColor(59,130,246),1.18)
card(s,3.72,"Pull-request CI","Require trusted policy and publish evidence where reviewers already work.",MINT,TEAL,1.18)
card(s,5.09,"Agent and platform workflows","Give coding agents, MCP hosts, AI-DLC, and policy engines one stable verification contract.",AMBER,RGBColor(245,158,11),1.35)
card(s,6.64,"Audit and calibration","Retain verdict history, anchor the chain head, and tune rules with reviewer labels.",ROSE,RGBColor(220,38,38),1.18)
txt(s,"The common outcome: less human time spent proving basic facts.",.75,8.35,6.5,.55,15,TEAL,True,PP_ALIGN.CENTER)

# 14 proof
s=base(14); heading(s,"What is verified today","Product evidence")
metrics=[("25","Maven reactor projects"),("844","tests"),("0","failures / errors"),("0","critical / high review findings")]
for i,(n,l) in enumerate(metrics):
    x=.55+(i%2)*3.55; y=2.35+(i//2)*1.75; shape(s,x,y,3.3,1.35,BLUE if i%2==0 else MINT,LINE)
    txt(s,n,x+.2,y+.22,.95,.5,28,TEAL,True); txt(s,l,x+1.15,y+.32,1.9,.62,12,NAVY,True)
bullets(s,["Full clean Maven verification and configured coverage gates pass.","AI-DLC adapter tests, preset parity, YAML, docs, visuals, and deck validators pass.","Independent adversarial review ended at Critical 0 / High 0.","One Windows capability-gated symlink test is skipped where the OS cannot create a link."],.7,6.15,6.6,2.25,13,INK)

# 15 adoption
s=base(15); heading(s,"Adopt without boiling the ocean","Four steps")
steps=[("1. Observe","Start with minimal local syntax + invariant."),("2. Calibrate","Add repository gates; keep structural signals advisory."),("3. Protect","Commit .aiv; require pinned/frozen trusted policy in CI."),("4. Enforce","Promote only objective, proven rules. Retain evidence.")]
for i,(h,b) in enumerate(steps):
    y=2.35+i*1.38; card(s,y,h,b,[BLUE,MINT,AMBER,ROSE][i],[RGBColor(59,130,246),TEAL,RGBColor(245,158,11),RGBColor(220,38,38)][i],1.15)
txt(s,"Enable the gates that solve your problem. Do not enable everything because it exists.",.75,8.25,6.5,.65,15,TEAL,True,PP_ALIGN.CENTER)

# 16 CTA
s=base(16,True)
txt(s,"Pilot one repository.",.6,1.0,6.8,.75,34,WHITE,True,PP_ALIGN.CENTER)
txt(s,"Measure one thing:",.6,2.05,6.8,.48,20,RGBColor(148,163,184),False,PP_ALIGN.CENTER)
txt(s,"How much reviewer time stops being spent proving basic facts?",.75,2.75,6.5,1.25,27,RGBColor(45,212,191),True,PP_ALIGN.CENTER)
shape(s,.75,4.65,6.5,2.35,RGBColor(20,31,61),RGBColor(51,65,85))
bullets(s,["Choose a real change workflow.","Run advisory first.","Protect policy in CI.","Review the evidence with maintainers.","Let measured value decide the next step."],1.05,4.98,5.9,1.75,14,WHITE)
txt(s,"github.com/vaquarkhan/aiv-integrity-gate",.7,7.72,6.6,.35,15,WHITE,True,PP_ALIGN.CENTER)
txt(s,"Viquar Khan  |  vaquar.khan@gmail.com",.7,8.40,6.6,.3,11,RGBColor(148,163,184),False,PP_ALIGN.CENTER)

prs.core_properties.title="AIV Integrity Gate: From Generated Code to Verifiable Change"
prs.core_properties.subject="Self-contained LinkedIn document carousel"
prs.core_properties.author="Viquar Khan"
prs.core_properties.keywords="AIV, AI-DLC, software integrity, coding agents, deterministic evidence"
prs.core_properties.comments="All claims grounded in the validated local AIV source tree."
prs.save(OUT)
print(f"WROTE {OUT} ({len(prs.slides)} pages, 4:5 portrait)")
